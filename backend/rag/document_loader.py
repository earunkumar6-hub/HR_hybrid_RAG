"""
Loads raw files (PDF / DOCX / TXT / MD) and splits them into overlapping
chunks ready for embedding into the DL (Document Layer) collection.
"""
from __future__ import annotations

import uuid
from pathlib import Path
from typing import List, Optional

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from config import settings

try:
    import docx  # python-docx
except ImportError:  # pragma: no cover - optional dependency
    docx = None


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    if docx is None:
        raise RuntimeError("python-docx is not installed; cannot read .docx files")
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".txt": _read_text,
    ".md": _read_text,
}

# Exposed so the folder-scan ingestor (rag/ingest.py) knows which files
# in backend/data/documents/ it can actually process.
SUPPORTED_SUFFIXES = set(_READERS.keys())


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    reader = _READERS.get(suffix)
    if reader is None:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_SUFFIXES)}"
        )
    return reader(path)


def load_and_chunk(
    path: Path, filename: str, doc_id: Optional[str] = None
) -> tuple[str, List[Document]]:
    """Extract text from `path` and split it into LangChain Documents.

    `doc_id` should be a *stable* identifier derived from the file's
    location (see rag/ingest.py) so re-running ingestion on an unchanged
    file upserts the same Chroma/graph records instead of duplicating
    them. If omitted, a random id is generated (fine for one-off/manual
    use, but callers that care about idempotency should always pass one).
    """
    doc_id = doc_id or uuid.uuid4().hex[:16]
    raw_text = extract_text(path)

    if not raw_text.strip():
        raise ValueError(
            f"No extractable text found in '{filename}'. "
            "If this is a scanned/image PDF it needs OCR before ingestion."
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    raw_chunks = splitter.split_text(raw_text)

    chunks = [
        Document(
            page_content=chunk,
            metadata={
                "doc_id": doc_id,
                "filename": filename,
                "chunk_index": idx,
            },
        )
        for idx, chunk in enumerate(raw_chunks)
    ]
    return doc_id, chunks
